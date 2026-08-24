"""Agent tools that GENERATE files (text/code/markdown, Word .docx, Excel .xlsx).

Each tool writes to the generated-files store and returns a human-readable string
with a hidden file marker appended (see genfiles.encode_file_marker), which the
agent loop turns into a downloadable/previewable file card in the UI.
"""

from __future__ import annotations

import csv
import io
from pathlib import Path

from backend.tools.genfiles import encode_file_marker, save_generated_file

# Extensions create_file may write (text-based only; no binaries/executables).
_TEXT_EXT = {
    ".md", ".markdown", ".txt", ".text", ".py", ".js", ".ts", ".jsx", ".tsx",
    ".json", ".csv", ".tsv", ".html", ".htm", ".css", ".svg", ".xml", ".yaml",
    ".yml", ".sql", ".sh", ".ini", ".toml", ".cfg", ".rst", ".log", ".c",
    ".cpp", ".h", ".hpp", ".rs", ".go", ".java", ".php", ".rb", ".swift", ".kt",
}


class FileGenError(ValueError):
    pass


def _finish(meta: dict, human: str) -> str:
    return f"{human}\n" + encode_file_marker(meta)


def create_file(filename: str, content: str) -> str:
    ext = Path(filename or "").suffix.lower()
    if not ext:
        raise FileGenError("filename phải có phần mở rộng, ví dụ ghi_chu.md")
    if ext not in _TEXT_EXT:
        raise FileGenError(f"đuôi '{ext}' không được phép (chỉ file text/code/markdown…)")
    data = (content or "").encode("utf-8")
    meta = save_generated_file(filename, data)
    return _finish(meta, f"Đã tạo file {meta['name']} ({meta['size']} bytes).")


def create_docx(filename: str, content: str) -> str:
    from docx import Document

    name = filename or "tai_lieu.docx"
    if not name.lower().endswith(".docx"):
        name += ".docx"
    doc = Document()
    for raw in (content or "").splitlines():
        line = raw.rstrip()
        stripped = line.strip()
        if not stripped:
            continue
        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            doc.add_heading(stripped.lstrip("#").strip(), level=min(max(level, 1), 4))
        elif stripped[:2] in ("- ", "* "):
            doc.add_paragraph(stripped[2:], style="List Bullet")
        else:
            doc.add_paragraph(line)
    buf = io.BytesIO()
    doc.save(buf)
    meta = save_generated_file(name, buf.getvalue())
    return _finish(meta, f"Đã tạo tài liệu Word {meta['name']}.")


def create_xlsx(filename: str, content: str) -> str:
    from openpyxl import Workbook

    name = filename or "bang_tinh.xlsx"
    if not name.lower().endswith(".xlsx"):
        name += ".xlsx"
    wb = Workbook()
    ws = wb.active
    rows = 0
    for row in csv.reader(io.StringIO(content or "")):
        ws.append(row)
        rows += 1
    if rows == 0:
        ws.append([])
    buf = io.BytesIO()
    wb.save(buf)
    meta = save_generated_file(name, buf.getvalue())
    return _finish(meta, f"Đã tạo bảng tính {meta['name']} ({rows} hàng).")
